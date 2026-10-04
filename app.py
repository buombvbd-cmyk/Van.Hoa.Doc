import os
from datetime import date, datetime, timedelta
from io import BytesIO
import re
import unicodedata
import base64

from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, abort
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
import qrcode

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'change-this-secret-key')

DATABASE_URL = os.environ.get('DATABASE_URL', '').strip()
if DATABASE_URL.startswith('postgres://'):
    DATABASE_URL = DATABASE_URL.replace('postgres://', 'postgresql://', 1)
if DATABASE_URL:
    app.config['SQLALCHEMY_DATABASE_URI'] = DATABASE_URL
else:
    app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///library.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024
UPLOAD_DIR = os.path.join(app.root_path, 'static', 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)
ALLOWED_COVER_EXTENSIONS = {'png','jpg','jpeg','webp'}

db = SQLAlchemy(app)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(30), default='admin')

class Book(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(80), unique=True, nullable=False)
    title = db.Column(db.String(255), nullable=False)
    author = db.Column(db.String(255))
    category = db.Column(db.String(120))
    publisher = db.Column(db.String(255))
    year = db.Column(db.Integer)
    quantity = db.Column(db.Integer, default=1)
    location = db.Column(db.String(120))
    cover = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Reader(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(80), unique=True, nullable=False)
    name = db.Column(db.String(255), nullable=False)
    class_name = db.Column(db.String(120))
    phone = db.Column(db.String(50))
    type = db.Column(db.String(50), default='Học sinh')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Loan(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    book_id = db.Column(db.Integer, db.ForeignKey('book.id'), nullable=False)
    reader_id = db.Column(db.Integer, db.ForeignKey('reader.id'), nullable=False)
    borrow_date = db.Column(db.Date, default=date.today)
    due_date = db.Column(db.Date, nullable=False)
    return_date = db.Column(db.Date)
    status = db.Column(db.String(30), default='Đang mượn')
    book = db.relationship('Book', backref='loans')
    reader = db.relationship('Reader', backref='loans')

class Shelf(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(80), unique=True, nullable=False)
    name = db.Column(db.String(255), nullable=False)
    zone = db.Column(db.String(120))
    note = db.Column(db.Text)

class ReadingActivity(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reader_id = db.Column(db.Integer, db.ForeignKey('reader.id'), nullable=False)
    activity_type = db.Column(db.String(80), nullable=False)
    title = db.Column(db.String(255), nullable=False)
    description = db.Column(db.Text)
    points = db.Column(db.Integer, default=0)
    status = db.Column(db.String(30), default='Đã duyệt')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    reader = db.relationship('Reader', backref='activities')

class ReadingPointsLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reader_id = db.Column(db.Integer, db.ForeignKey('reader.id'), nullable=False)
    loan_id = db.Column(db.Integer, db.ForeignKey('loan.id'), unique=True)
    points = db.Column(db.Integer, default=20)
    reason = db.Column(db.String(255))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Badge(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reader_id = db.Column(db.Integer, db.ForeignKey('reader.id'), nullable=False)
    code = db.Column(db.String(80), nullable=False)
    name = db.Column(db.String(120), nullable=False)
    earned_at = db.Column(db.DateTime, default=datetime.utcnow)
    __table_args__ = (db.UniqueConstraint('reader_id', 'code', name='uq_reader_badge'),)

class Reward(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reader_id = db.Column(db.Integer, db.ForeignKey('reader.id'), nullable=False)
    title = db.Column(db.String(255), nullable=False)
    description = db.Column(db.Text)
    points = db.Column(db.Integer, default=0)
    awarded_at = db.Column(db.DateTime, default=datetime.utcnow)
    reader = db.relationship('Reader', backref='rewards')

BADGES = [(1, 'Người đọc mới'), (5, 'Bạn đọc chăm chỉ'), (10, 'Mọt sách'), (20, 'Người đọc tích cực'), (50, 'Đại sứ Văn hóa Đọc')]
ACTIVITY_POINTS = {'Cảm nhận sách': 10, 'Giới thiệu sách': 20, 'Video giới thiệu sách': 30, 'Sân khấu hóa': 30, 'Khác': 5}

@app.context_processor
def inject_globals():
    return {'today': date.today(), 'app_name': 'QUẢN LÍ THƯ VIỆN'}

def logged_in():
    return bool(session.get('user'))

def require_login():
    if not logged_in():
        return redirect(url_for('login'))
    return None

def total_reading_points(reader_id):
    return ((db.session.query(db.func.coalesce(db.func.sum(ReadingPointsLog.points), 0)).filter_by(reader_id=reader_id).scalar() or 0)
            + (db.session.query(db.func.coalesce(db.func.sum(ReadingActivity.points), 0)).filter_by(reader_id=reader_id, status='Đã duyệt').scalar() or 0)
            + (db.session.query(db.func.coalesce(db.func.sum(Reward.points), 0)).filter_by(reader_id=reader_id).scalar() or 0))

def books_read(reader_id):
    return Loan.query.filter_by(reader_id=reader_id, status='Đã trả').count()

def refresh_badges(reader_id):
    count = books_read(reader_id)
    for threshold, name in BADGES:
        if count >= threshold and not Badge.query.filter_by(reader_id=reader_id, code=f'b{threshold}').first():
            db.session.add(Badge(reader_id=reader_id, code=f'b{threshold}', name=name))
    db.session.commit()

def award_return_points(loan):
    if not ReadingPointsLog.query.filter_by(loan_id=loan.id).first():
        db.session.add(ReadingPointsLog(reader_id=loan.reader_id, loan_id=loan.id, points=20, reason='Hoàn trả sách'))
        db.session.commit()
    refresh_badges(loan.reader_id)

def is_overdue(loan):
    return loan.status == 'Đang mượn' and loan.due_date and loan.due_date < date.today()

def class_names():
    return [x[0] for x in db.session.query(Reader.class_name).filter(Reader.class_name.isnot(None), Reader.class_name!='').distinct().order_by(Reader.class_name).all()]

def init_db():
    db.create_all()
    if not User.query.filter_by(username='admin').first():
        db.session.add(User(username='admin', password_hash=generate_password_hash('admin123'), role='admin'))
        db.session.commit()

with app.app_context():
    init_db()
    # SQLAlchemy create_all() không đổi kiểu cột đã tồn tại.
    # PostgreSQL cần chuyển cover sang TEXT để chứa data URL của ảnh.
    try:
        if DATABASE_URL:
            db.session.execute(db.text('ALTER TABLE book ALTER COLUMN cover TYPE TEXT'))
            db.session.commit()
    except Exception:
        db.session.rollback()

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username','').strip()
        password = request.form.get('password','')
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password_hash, password):
            session['user'] = user.username
            session['role'] = user.role
            return redirect(url_for('index'))
        flash('Tài khoản hoặc mật khẩu không đúng.', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/')
def index():
    if not logged_in(): return redirect(url_for('login'))
    q = request.args.get('q','').strip()
    query = Book.query
    if q:
        like=f'%{q}%'
        query=query.filter(db.or_(Book.code.ilike(like), Book.title.ilike(like), Book.author.ilike(like), Book.category.ilike(like)))
    books=query.order_by(Book.id.desc()).all()
    stats={
        'total_books': sum(max(0,b.quantity) for b in Book.query.all()),
        'titles': Book.query.count(), 'readers': Reader.query.count(),
        'borrowing': Loan.query.filter_by(status='Đang mượn').count(),
        'overdue': Loan.query.filter(Loan.status=='Đang mượn', Loan.due_date < date.today()).count(),
        'activities': ReadingActivity.query.count()
    }
    return render_template('index.html', books=books, stats=stats, q=q)

@app.route('/books')
def books():
    if not logged_in(): return redirect(url_for('login'))
    q=request.args.get('q','').strip(); query=Book.query
    if q:
        like=f'%{q}%'; query=query.filter(db.or_(Book.code.ilike(like),Book.title.ilike(like),Book.author.ilike(like)))
    return render_template('books.html', books=query.order_by(Book.title).all(), q=q)

def save_cover(file_obj):
    """Lưu ảnh bìa trực tiếp trong database để không mất trên Render."""
    if not file_obj or not file_obj.filename:
        return None

    ext = file_obj.filename.rsplit('.', 1)[-1].lower() if '.' in file_obj.filename else ''
    if ext not in ALLOWED_COVER_EXTENSIONS:
        raise ValueError('Ảnh bìa phải là PNG, JPG, JPEG hoặc WEBP.')

    raw = file_obj.read()
    if not raw:
        raise ValueError('Ảnh bìa không có dữ liệu.')

    mime = 'image/jpeg' if ext in {'jpg', 'jpeg'} else f'image/{ext}'
    encoded = base64.b64encode(raw).decode('ascii')
    return f'data:{mime};base64,{encoded}'

@app.route('/books/add', methods=['GET','POST'])
def add_book():
    if not logged_in(): return redirect(url_for('login'))
    if request.method=='POST':
        try:
            code=request.form.get('code','').strip(); title=request.form.get('title','').strip()
            if not code or not title: raise ValueError('Mã sách và tên sách là bắt buộc.')
            if Book.query.filter_by(code=code).first(): raise ValueError('Mã sách đã tồn tại.')
            cover=save_cover(request.files.get('cover'))
            b=Book(code=code,title=title,author=request.form.get('author'),category=request.form.get('category'),publisher=request.form.get('publisher'),year=int(request.form['year']) if request.form.get('year') else None,quantity=max(0,int(request.form.get('quantity') or 0)),location=request.form.get('location'),cover=cover)
            db.session.add(b); db.session.commit(); flash('Đã thêm sách.','success'); return redirect(url_for('books'))
        except Exception as e:
            db.session.rollback(); flash(f'Không thể thêm sách: {e}','error')
    return render_template('book_form.html', book=None)

@app.route('/books/edit/<int:book_id>', methods=['GET','POST'])
def edit_book(book_id):
    if not logged_in(): return redirect(url_for('login'))
    b=Book.query.get_or_404(book_id)
    if request.method=='POST':
        try:
            code=request.form.get('code','').strip(); title=request.form.get('title','').strip()
            other=Book.query.filter(Book.code==code,Book.id!=b.id).first()
            if other: raise ValueError('Mã sách đã tồn tại.')
            b.code=code; b.title=title; b.author=request.form.get('author'); b.category=request.form.get('category'); b.publisher=request.form.get('publisher'); b.year=int(request.form['year']) if request.form.get('year') else None; b.quantity=max(0,int(request.form.get('quantity') or 0)); b.location=request.form.get('location')
            new_cover=save_cover(request.files.get('cover'))
            if new_cover: b.cover=new_cover
            db.session.commit(); flash('Đã cập nhật sách.','success'); return redirect(url_for('book_detail',book_id=b.id))
        except Exception as e: db.session.rollback(); flash(f'Lỗi: {e}','error')
    return render_template('book_form.html', book=b)

@app.route('/books/delete/<int:book_id>', methods=['POST'])
def delete_book(book_id):
    if not logged_in(): return redirect(url_for('login'))
    b=Book.query.get_or_404(book_id)
    if Loan.query.filter_by(book_id=b.id).first(): flash('Không thể xóa sách đã có lịch sử mượn.','error')
    else:
        db.session.delete(b); db.session.commit(); flash('Đã xóa sách.','success')
    return redirect(url_for('books'))

@app.route('/books/<int:book_id>')
def book_detail(book_id):
    if not logged_in(): return redirect(url_for('login'))
    b=Book.query.get_or_404(book_id); return render_template('book_detail.html', book=b)

@app.route('/books/<int:book_id>/qr')
def book_qr(book_id):
    if not logged_in(): return redirect(url_for('login'))
    b=Book.query.get_or_404(book_id); img=qrcode.make(f'BOOK:{b.code}'); out=BytesIO(); img.save(out,format='PNG'); out.seek(0); return send_file(out,mimetype='image/png')

@app.route('/books/export')
def export_books():
    if not logged_in(): return redirect(url_for('login'))
    wb=Workbook(); ws=wb.active; ws.title='Danh sách sách'; ws.append(['Mã','Tên sách','Tác giả','Thể loại','Nhà xuất bản','Năm','Số lượng','Vị trí'])
    for b in Book.query.order_by(Book.id).all(): ws.append([b.code,b.title,b.author,b.category,b.publisher,b.year,b.quantity,b.location])
    out=BytesIO(); wb.save(out); out.seek(0); return send_file(out,as_attachment=True,download_name='danh_sach_sach.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

def _norm_header(value):
    """Chuẩn hóa tiêu đề Excel để nhận nhiều cách đặt tên khác nhau."""
    if value is None:
        return ''
    text = str(value).strip().lower()
    text = ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')
    text = text.replace('đ', 'd')
    return re.sub(r'[^a-z0-9]+', '', text)

EXCEL_HEADER_ALIASES = {
    'code': {'ma', 'masach', 'masachcode', 'code', 'bookcode', 'isbn'},
    'title': {'tensach', 'tensachbook', 'tuaasach', 'title', 'booktitle'},
    'author': {'tacgia', 'author', 'bookauthor'},
    'category': {'theloai', 'loaisach', 'category', 'genre'},
    'publisher': {'nhaxuatban', 'nxb', 'publisher'},
    'year': {'nam', 'namxuatban', 'year', 'publicationyear'},
    'quantity': {'soluong', 'soquyen', 'quantity', 'qty', 'tonkho'},
    'location': {'vitri', 'kesach', 'kệsach', 'location', 'shelf', 'shelflocation'},
}

def _excel_column_map(headers):
    mapping = {}
    normalized = [_norm_header(h) for h in headers]
    for idx, h in enumerate(normalized):
        for field, aliases in EXCEL_HEADER_ALIASES.items():
            if h in aliases and field not in mapping:
                mapping[field] = idx
    return mapping

def _cell(row, mapping, field, default=None):
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return default
    value = row[idx]
    if value is None:
        return default
    if isinstance(value, str):
        value = value.strip()
        return value if value else default
    return value

def _to_int(value, default=None):
    if value is None or value == '':
        return default
    try:
        return int(float(str(value).replace(',', '.')))
    except (TypeError, ValueError):
        return default

@app.route('/books/import-template')
def import_books_template():
    if not logged_in(): return redirect(url_for('login'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Danh sách sách'
    ws.append(['Mã', 'Tên sách', 'Tác giả', 'Thể loại', 'Nhà xuất bản', 'Năm', 'Số lượng', 'Vị trí'])
    ws.append(['S001', 'Ví dụ: Dế Mèn phiêu lưu ký', 'Tô Hoài', 'Văn học', 'Kim Đồng', 2026, 5, 'Kệ A1'])
    ws.freeze_panes = 'A2'
    widths = [16, 35, 28, 22, 28, 12, 14, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width
    out = BytesIO(); wb.save(out); out.seek(0)
    return send_file(out, as_attachment=True, download_name='mau_nhap_sach.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/books/cleanup-bad-imports', methods=['POST'])
def cleanup_bad_imports():
    if not logged_in():
        return redirect(url_for('login'))
    # Chỉ nhắm đúng nhóm dữ liệu nhập lỗi đã xác định: tác giả=4, thể loại=44, số lượng=44.
    # Không xóa sách đã có lịch sử mượn.
    candidates = Book.query.filter_by(author='4', category='44', quantity=44).all()
    removed = 0
    protected = 0
    for b in candidates:
        if Loan.query.filter_by(book_id=b.id).first():
            protected += 1
            continue
        db.session.delete(b)
        removed += 1
    db.session.commit()
    flash(f'Đã dọn {removed} sách nhập lỗi. Giữ lại {protected} sách vì đã có lịch sử mượn.', 'success')
    return redirect(url_for('import_books'))

@app.route('/books/import', methods=['GET','POST'])
def import_books():
    if not logged_in(): return redirect(url_for('login'))
    if request.method == 'POST':
        f = request.files.get('file')
        if not f:
            flash('Chưa chọn file Excel.', 'error')
            return redirect(request.url)
        try:
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                flash('File Excel không có dữ liệu.', 'error')
                return redirect(request.url)

            headers = list(rows[0])
            mapping = _excel_column_map(headers)
            required = {'code', 'title'}
            missing = required - set(mapping)

            # Không tự đoán thứ tự cột nữa: tránh nhập lệch dữ liệu.
            if missing:
                missing_text = ', '.join(sorted(missing))
                flash(f'File Excel chưa đúng mẫu. Thiếu cột bắt buộc: {missing_text}. Hãy tải file Excel mẫu rồi nhập lại.', 'error')
                return redirect(url_for('import_books'))
            start_row = 1

            added = 0
            updated = 0
            skipped = 0
            for row in rows[start_row:]:
                code = _cell(row, mapping, 'code')
                title = _cell(row, mapping, 'title')
                if not code or not title:
                    skipped += 1
                    continue
                code = str(code).strip()
                title = str(title).strip()
                data = {
                    'title': title,
                    'author': _cell(row, mapping, 'author'),
                    'category': _cell(row, mapping, 'category'),
                    'publisher': _cell(row, mapping, 'publisher'),
                    'year': _to_int(_cell(row, mapping, 'year')),
                    'quantity': _to_int(_cell(row, mapping, 'quantity'), 0),
                    'location': _cell(row, mapping, 'location'),
                }
                b = Book.query.filter_by(code=code).first()
                if b:
                    for key, value in data.items():
                        setattr(b, key, value)
                    updated += 1
                else:
                    db.session.add(Book(code=code, **data))
                    added += 1
            db.session.commit()
            flash(f'Đã nhập Excel: thêm {added}, cập nhật {updated}, bỏ qua {skipped} dòng.', 'success')
            return redirect(url_for('books'))
        except Exception as e:
            db.session.rollback()
            flash(f'Không thể đọc Excel: {e}', 'error')
    return render_template('import_books.html')

@app.route('/readers')
def readers():
    if not logged_in(): return redirect(url_for('login'))
    q=request.args.get('q','').strip(); query=Reader.query
    if q:
        like=f'%{q}%'; query=query.filter(db.or_(Reader.code.ilike(like),Reader.name.ilike(like),Reader.class_name.ilike(like)))
    data=query.order_by(Reader.name).all()
    rows=[{'reader':r,'books':books_read(r.id),'points':total_reading_points(r.id)} for r in data]
    return render_template('readers.html', rows=rows, q=q)

@app.route('/readers/add', methods=['GET','POST'])
def add_reader():
    if not logged_in(): return redirect(url_for('login'))
    if request.method=='POST':
        try: db.session.add(Reader(code=request.form['code'].strip(),name=request.form['name'].strip(),class_name=request.form.get('class_name'),phone=request.form.get('phone'),type=request.form.get('type') or 'Học sinh')); db.session.commit(); flash('Đã thêm bạn đọc.','success'); return redirect(url_for('readers'))
        except Exception as e: db.session.rollback(); flash(f'Lỗi: {e}','error')
    return render_template('reader_form.html',reader=None)

@app.route('/readers/edit/<int:reader_id>', methods=['GET','POST'])
def edit_reader(reader_id):
    if not logged_in(): return redirect(url_for('login'))
    r=Reader.query.get_or_404(reader_id)
    if request.method=='POST':
        try: r.code=request.form['code'].strip(); r.name=request.form['name'].strip(); r.class_name=request.form.get('class_name'); r.phone=request.form.get('phone'); r.type=request.form.get('type') or 'Học sinh'; db.session.commit(); flash('Đã cập nhật bạn đọc.','success'); return redirect(url_for('readers'))
        except Exception as e: db.session.rollback(); flash(f'Lỗi: {e}','error')
    return render_template('reader_form.html',reader=r)

@app.route('/readers/delete/<int:reader_id>', methods=['POST'])
def delete_reader(reader_id):
    if not logged_in(): return redirect(url_for('login'))
    r=Reader.query.get_or_404(reader_id)
    if Loan.query.filter_by(reader_id=r.id).first(): flash('Không thể xóa bạn đọc đã có lịch sử mượn.','error')
    else: db.session.delete(r); db.session.commit(); flash('Đã xóa bạn đọc.','success')
    return redirect(url_for('readers'))

@app.route('/readers/<int:reader_id>')
def reader_detail(reader_id):
    if not logged_in(): return redirect(url_for('login'))
    r=Reader.query.get_or_404(reader_id); return render_template('reader_detail.html',reader=r,loans=Loan.query.filter_by(reader_id=r.id).order_by(Loan.id.desc()).all())

@app.route('/readers/<int:reader_id>/qr')
def reader_qr(reader_id):
    if not logged_in(): return redirect(url_for('login'))
    r=Reader.query.get_or_404(reader_id); img=qrcode.make(f'READER:{r.code}'); out=BytesIO(); img.save(out,format='PNG'); out.seek(0); return send_file(out,mimetype='image/png')

@app.route('/readers/<int:reader_id>/reading-profile')
def reading_profile(reader_id):
    if not logged_in(): return redirect(url_for('login'))
    r=Reader.query.get_or_404(reader_id); refresh_badges(r.id); return render_template('reading_profile.html',reader=r,books_read=books_read(r.id),points=total_reading_points(r.id),badges=Badge.query.filter_by(reader_id=r.id).order_by(Badge.earned_at).all(),activities=ReadingActivity.query.filter_by(reader_id=r.id).order_by(ReadingActivity.id.desc()).all(),rewards=Reward.query.filter_by(reader_id=r.id).order_by(Reward.id.desc()).all(),loans=Loan.query.filter_by(reader_id=r.id).order_by(Loan.id.desc()).all())

@app.route('/readers/<int:reader_id>/card')
def reader_card(reader_id):
    if not logged_in(): return redirect(url_for('login'))
    return render_template('reader_card.html',reader=Reader.query.get_or_404(reader_id))

@app.route('/loans')
def loans():
    if not logged_in(): return redirect(url_for('login'))
    q=request.args.get('q','').strip(); status=request.args.get('status','all')
    query=Loan.query.join(Book).join(Reader)
    if q:
        like=f'%{q}%'; query=query.filter(db.or_(Book.title.ilike(like),Book.code.ilike(like),Reader.name.ilike(like),Reader.code.ilike(like)))
    if status == 'active': query=query.filter(Loan.status=='Đang mượn')
    elif status == 'returned': query=query.filter(Loan.status=='Đã trả')
    elif status == 'overdue': query=query.filter(Loan.status=='Đang mượn', Loan.due_date < date.today())
    data=query.order_by(Loan.id.desc()).all()
    return render_template('loans.html',loans=data,q=q,status=status,is_overdue=is_overdue)

@app.route('/loans/add', methods=['GET','POST'])
def add_loan():
    if not logged_in(): return redirect(url_for('login'))
    if request.method=='POST':
        try:
            book=Book.query.get(int(request.form['book_id'])); reader=Reader.query.get(int(request.form['reader_id']))
            if not book or not reader: raise ValueError('Sách hoặc bạn đọc không tồn tại.')
            if book.quantity <= 0: raise ValueError('Sách đã hết.')
            if Loan.query.filter_by(reader_id=reader.id,status='Đang mượn').count() >= 5: raise ValueError('Bạn đọc đã đạt tối đa 5 sách đang mượn.')
            due=date.today()+timedelta(days=int(request.form.get('days') or 14)); db.session.add(Loan(book_id=book.id,reader_id=reader.id,borrow_date=date.today(),due_date=due)); book.quantity-=1; db.session.commit(); flash('Đã lập phiếu mượn.','success'); return redirect(url_for('loans'))
        except Exception as e: db.session.rollback(); flash(str(e),'error')
    return render_template('loan_form.html',books=Book.query.filter(Book.quantity>0).order_by(Book.title).all(),readers=Reader.query.order_by(Reader.name).all())

@app.route('/loans/return/<int:loan_id>', methods=['POST'])
def return_loan(loan_id):
    if not logged_in(): return redirect(url_for('login'))
    l=Loan.query.get_or_404(loan_id)
    if l.status=='Đã trả': flash('Phiếu này đã trả.','error'); return redirect(url_for('loans'))
    l.status='Đã trả'; l.return_date=date.today(); l.book.quantity+=1; db.session.commit(); award_return_points(l); flash('Đã trả sách và cộng 20 điểm Văn hóa Đọc.','success'); return redirect(url_for('loans'))

@app.route('/quick-borrow', methods=['GET','POST'])
def quick_borrow():
    if not logged_in(): return redirect(url_for('login'))
    if request.method=='POST':
        code=request.form.get('code','').strip(); book=Book.query.filter_by(code=code).first(); reader=Reader.query.filter_by(code=request.form.get('reader_code','').strip()).first()
        if not book or not reader: flash('Không tìm thấy mã sách hoặc mã bạn đọc.','error')
        elif book.quantity<=0: flash('Sách đã hết.','error')
        elif Loan.query.filter_by(reader_id=reader.id,status='Đang mượn').count()>=5: flash('Bạn đọc đã đạt tối đa 5 sách.','error')
        else:
            db.session.add(Loan(book_id=book.id,reader_id=reader.id,borrow_date=date.today(),due_date=date.today()+timedelta(days=14)))
            book.quantity-=1
            db.session.commit()
            flash('Mượn nhanh thành công.','success')
    return render_template('quick_borrow.html')

@app.route('/quick-return', methods=['GET','POST'])
def quick_return():
    if not logged_in(): return redirect(url_for('login'))
    if request.method=='POST':
        code=request.form.get('code','').strip(); reader_code=request.form.get('reader_code','').strip(); q=Loan.query.join(Book).join(Reader).filter(Loan.status=='Đang mượn')
        if code: q=q.filter(Book.code==code)
        if reader_code: q=q.filter(Reader.code==reader_code)
        l=q.order_by(Loan.id.asc()).first()
        if not l: flash('Không tìm thấy phiếu mượn đang hoạt động.','error')
        else: l.status='Đã trả'; l.return_date=date.today(); l.book.quantity+=1; db.session.commit(); award_return_points(l); flash('Trả nhanh thành công và cộng 20 điểm.','success')
    return render_template('quick_return.html')

@app.route('/shelves')
def shelves():
    if not logged_in(): return redirect(url_for('login'))
    return render_template('shelves.html',shelves=Shelf.query.order_by(Shelf.code).all())

@app.route('/shelves/add', methods=['GET','POST'])
def add_shelf():
    if not logged_in(): return redirect(url_for('login'))
    if request.method=='POST':
        try: db.session.add(Shelf(code=request.form['code'].strip(),name=request.form['name'].strip(),zone=request.form.get('zone'),note=request.form.get('note'))); db.session.commit(); flash('Đã thêm kệ.','success'); return redirect(url_for('shelves'))
        except Exception as e: db.session.rollback(); flash(f'Lỗi: {e}','error')
    return render_template('shelf_form.html',shelf=None)

@app.route('/shelves/edit/<int:shelf_id>', methods=['GET','POST'])
def edit_shelf(shelf_id):
    if not logged_in(): return redirect(url_for('login'))
    s=Shelf.query.get_or_404(shelf_id)
    if request.method=='POST':
        try: s.code=request.form['code'].strip(); s.name=request.form['name'].strip(); s.zone=request.form.get('zone'); s.note=request.form.get('note'); db.session.commit(); flash('Đã cập nhật kệ.','success'); return redirect(url_for('shelves'))
        except Exception as e: db.session.rollback(); flash(f'Lỗi: {e}','error')
    return render_template('shelf_form.html',shelf=s)

@app.route('/shelves/delete/<int:shelf_id>', methods=['POST'])
def delete_shelf(shelf_id):
    if not logged_in(): return redirect(url_for('login'))
    s=Shelf.query.get_or_404(shelf_id); db.session.delete(s); db.session.commit(); flash('Đã xóa kệ.','success'); return redirect(url_for('shelves'))

@app.route('/classes')
def classes():
    if not logged_in(): return redirect(url_for('login'))
    names=[x[0] for x in db.session.query(Reader.class_name).filter(Reader.class_name.isnot(None),Reader.class_name!='').distinct().order_by(Reader.class_name).all()]
    data=[]
    for name in names:
        rs=Reader.query.filter_by(class_name=name).all(); data.append({'name':name,'readers':len(rs),'books':sum(books_read(r.id) for r in rs),'points':sum(total_reading_points(r.id) for r in rs)})
    return render_template('classes.html',classes=data)

@app.route('/reports')
def reports():
    if not logged_in(): return redirect(url_for('login'))
    return render_template('reports.html',total_books=Book.query.count(),copies=sum(b.quantity for b in Book.query.all()),readers=Reader.query.count(),loans=Loan.query.count(),active=Loan.query.filter_by(status='Đang mượn').count(),overdue=Loan.query.filter(Loan.status=='Đang mượn',Loan.due_date<date.today()).count(),points=db.session.query(db.func.coalesce(db.func.sum(ReadingPointsLog.points),0)).scalar() or 0)

@app.route('/reading-culture')
def reading_culture():
    if not logged_in(): return redirect(url_for('login'))
    cls=request.args.get('class','').strip(); query=Reader.query
    if cls: query=query.filter_by(class_name=cls)
    rows=[]
    for r in query.all():
        rows.append({'reader':r,'books':books_read(r.id),'points':total_reading_points(r.id),'badges':Badge.query.filter_by(reader_id=r.id).count()})
    rows.sort(key=lambda x:(-x['points'],-x['books'],x['reader'].name.lower()))
    for i,row in enumerate(rows,1): row['rank']=i
    return render_template('reading_culture.html',rows=rows,classes=class_names(),selected_class=cls)

@app.route('/reading/activities/add', methods=['POST'])
def add_activity():
    if not logged_in(): return redirect(url_for('reading_culture'))
    reader=Reader.query.get_or_404(int(request.form['reader_id']))
    typ=request.form.get('activity_type') or 'Khác'
    pts=max(0,int(request.form.get('points') or ACTIVITY_POINTS.get(typ,5)))
    status='Đã duyệt' if session.get('role')=='admin' else 'Chờ duyệt'
    db.session.add(ReadingActivity(reader_id=reader.id,activity_type=typ,title=request.form['title'].strip(),description=request.form.get('description'),points=pts,status=status))
    db.session.commit()
    flash('Đã ghi nhận hoạt động Văn hóa Đọc.' + (' Chờ duyệt.' if status=='Chờ duyệt' else ''),'success')
    return redirect(url_for('reading_profile',reader_id=reader.id))

@app.route('/reading/activities/<int:activity_id>/approve', methods=['POST'])
def approve_activity(activity_id):
    if not logged_in() or session.get('role')!='admin': abort(403)
    a=ReadingActivity.query.get_or_404(activity_id); a.status='Đã duyệt'; db.session.commit(); flash('Đã duyệt hoạt động.','success'); return redirect(url_for('reading_profile',reader_id=a.reader_id))

@app.route('/reading/activities/<int:activity_id>/reject', methods=['POST'])
def reject_activity(activity_id):
    if not logged_in() or session.get('role')!='admin': abort(403)
    a=ReadingActivity.query.get_or_404(activity_id); a.status='Từ chối'; db.session.commit(); flash('Đã từ chối hoạt động.','success'); return redirect(url_for('reading_profile',reader_id=a.reader_id))

@app.route('/reading/rewards/add', methods=['POST'])
def add_reward():
    if not logged_in(): return redirect(url_for('reading_culture'))
    reader=Reader.query.get_or_404(int(request.form['reader_id'])); db.session.add(Reward(reader_id=reader.id,title=request.form['title'].strip(),description=request.form.get('description'),points=int(request.form.get('points') or 0))); db.session.commit(); flash('Đã ghi nhận khen thưởng.','success'); return redirect(url_for('reading_profile',reader_id=reader.id))


@app.route('/reports/export')
def reports_export():
    if not logged_in(): return redirect(url_for('login'))
    wb=Workbook(); ws=wb.active; ws.title='Tong quan'
    ws.append(['Chỉ tiêu','Giá trị'])
    ws.append(['Đầu sách',Book.query.count()]); ws.append(['Tổng bản sách',sum(b.quantity for b in Book.query.all())])
    ws.append(['Bạn đọc',Reader.query.count()]); ws.append(['Phiếu mượn',Loan.query.count()]); ws.append(['Đang mượn',Loan.query.filter_by(status='Đang mượn').count()])
    ws.append(['Quá hạn',Loan.query.filter(Loan.status=='Đang mượn',Loan.due_date<date.today()).count()])
    ws.append(['Điểm Văn hóa Đọc',sum(total_reading_points(r.id) for r in Reader.query.all())])
    out=BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name='bao_cao_thu_vien.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/classes/export')
def classes_export():
    if not logged_in(): return redirect(url_for('login'))
    names=[x[0] for x in db.session.query(Reader.class_name).filter(Reader.class_name.isnot(None),Reader.class_name!='').distinct().order_by(Reader.class_name).all()]
    wb=Workbook(); ws=wb.active; ws.title='Theo lop'; ws.append(['Lớp','Bạn đọc','Lượt đọc','Điểm'])
    for name in names:
        rs=Reader.query.filter_by(class_name=name).all(); ws.append([name,len(rs),sum(books_read(r.id) for r in rs),sum(total_reading_points(r.id) for r in rs)])
    out=BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name='bao_cao_lop.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/books/<int:book_id>/history')
def book_history(book_id):
    if not logged_in(): return redirect(url_for('login'))
    b=Book.query.get_or_404(book_id)
    return render_template('book_history.html',book=b,loans=Loan.query.filter_by(book_id=b.id).order_by(Loan.id.desc()).all())

if __name__ == '__main__':
    print('Website QUẢN LÍ THƯ VIỆN đang chạy')
    app.run(host='0.0.0.0',port=int(os.environ.get('PORT',5000)),debug=False)
